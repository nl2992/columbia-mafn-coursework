#!/usr/bin/env python3
"""Stages 1-3 for the local Course Archive RAG.

The pipeline is intentionally source-preserving:

    python scripts/rag_pipeline.py inventory
    python scripts/rag_pipeline.py extract
    python scripts/rag_pipeline.py normalize

Generated files are written below .rag/ and are rebuildable. Source files are
never modified.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator


ROOT = Path(__file__).resolve().parents[1]
RAG_DIR = ROOT / ".rag"
EXTRACTED_DIR = RAG_DIR / "extracted"

IGNORED_DIRS = {".git", ".rag", ".DS_Store", "__pycache__", ".ipynb_checkpoints"}
CONTENT_ROLES = {
    "lectures",
    "assignments",
    "assessments",
    "solutions",
    "readings",
    "reference",
    "resources",
    "computational",
    "code",
    "data",
    "workbooks",
    "experiments",
    "source",
    "sessions",
    "topics",
    "applications",
    "syllabi",
    "assets",
}
SUPPORTED_ROUTES = {
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".ppt": "legacy_presentation",
    ".docx": "docx",
    ".doc": "legacy_document",
    ".xlsx": "spreadsheet",
    ".xlsm": "spreadsheet",
    ".xls": "legacy_spreadsheet",
    ".csv": "dataset",
    ".tsv": "dataset",
    ".ipynb": "notebook",
    ".py": "source_code",
    ".c": "source_code",
    ".m": "source_code",
    ".tex": "source_code",
    ".txt": "text",
    ".md": "markdown",
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
    ".gif": "image",
    ".zip": "archive",
}
LARGE_DATA_BYTES = 5 * 1024 * 1024
TEXT_CHUNK_CHARS = 3500
TEXT_CHUNK_OVERLAP = 400


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def json_dump(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def jsonl_write(path: Path, rows: Iterable[dict[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
            count += 1
    return count


def jsonl_read(path: Path) -> Iterator[dict[str, Any]]:
    if not path.exists():
        return
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            try:
                yield json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSONL at {path}:{line_number}: {exc}") from exc


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            block = handle.read(1024 * 1024)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def relative_path(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def is_ignored(path: Path) -> bool:
    return any(part in IGNORED_DIRS for part in path.parts)


def infer_context(rel: str) -> dict[str, Any]:
    parts = Path(rel).parts
    term = parts[0] if parts and re.fullmatch(r"(?:Fall|Spring|Summer|Winter) \d{4}", parts[0]) else None
    if term:
        course = parts[1] if len(parts) > 1 else None
        context_start = 2
    elif parts and parts[0] == "Program-wide":
        course = None
        context_start = 1
    else:
        course = None
        context_start = 0

    role = None
    for segment in parts[context_start:-1]:
        if segment.lower() in CONTENT_ROLES:
            role = segment.lower()
            break
    if role is None and len(parts) > context_start + 1:
        role = parts[context_start + 1].lower()

    access_review = "not_reviewed"
    lowered = {segment.lower() for segment in parts}
    if "solutions" in lowered or "assessments" in lowered:
        access_review = "review_required"

    return {
        "term": term,
        "course": course,
        "content_role": role,
        "access_review": access_review,
    }


def route_for(path: Path) -> str:
    return SUPPORTED_ROUTES.get(path.suffix.lower(), "unsupported")


def review_reasons(record: dict[str, Any]) -> list[str]:
    reasons: list[str] = []
    if record["route"] == "unsupported":
        reasons.append("unsupported_extension")
    if record["route"].startswith("legacy_"):
        reasons.append("legacy_office_format")
    if record["route"] == "dataset" and record["size_bytes"] >= LARGE_DATA_BYTES:
        reasons.append("large_structured_dataset")
    if record["access_review"] == "review_required":
        reasons.append("assessment_or_solution_material")
    if record["route"] == "image":
        reasons.append("visual_asset_needs_ocr_or_caption_review")
    return reasons


def command_output(args: list[str], timeout: int = 600) -> subprocess.CompletedProcess[str]:
    try:
        return subprocess.run(args, capture_output=True, text=True, timeout=timeout, check=False)
    except FileNotFoundError as exc:
        raise RuntimeError(f"Required command not found: {args[0]}") from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"Command timed out after {timeout}s: {' '.join(args)}") from exc


def inventory() -> None:
    RAG_DIR.mkdir(parents=True, exist_ok=True)
    paths = sorted(
        (path for path in ROOT.rglob("*") if path.is_file() and not is_ignored(path)),
        key=lambda path: relative_path(path).lower(),
    )
    records: list[dict[str, Any]] = []
    by_hash: defaultdict[str, list[str]] = defaultdict(list)
    total_bytes = 0
    for index, path in enumerate(paths, 1):
        rel = relative_path(path)
        stat = path.stat()
        content_hash = sha256_file(path)
        context = infer_context(rel)
        record: dict[str, Any] = {
            "source_path": rel,
            "source_name": path.name,
            "extension": path.suffix.lower(),
            "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
            "size_bytes": stat.st_size,
            "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            "content_hash": content_hash,
            "document_id": f"sha256:{content_hash}",
            "route": route_for(path),
            "ingestion_status": "pending",
            "duplicate_of": None,
            **context,
        }
        record["review_reasons"] = review_reasons(record)
        records.append(record)
        by_hash[content_hash].append(rel)
        total_bytes += stat.st_size
        if index % 50 == 0 or index == len(paths):
            print(f"inventory: {index}/{len(paths)} files", file=sys.stderr)

    for record in records:
        paths_for_hash = by_hash[record["content_hash"]]
        record["source_paths"] = paths_for_hash
        if len(paths_for_hash) > 1 and record["source_path"] != paths_for_hash[0]:
            record["duplicate_of"] = paths_for_hash[0]

    jsonl_write(RAG_DIR / "manifest.jsonl", records)
    review_rows = (
        {
            "source_path": record["source_path"],
            "document_id": record["document_id"],
            "review_reasons": record["review_reasons"],
            "status": "pending",
        }
        for record in records
        if record["review_reasons"]
    )
    review_count = jsonl_write(RAG_DIR / "review-queue.jsonl", review_rows)

    route_counts = Counter(record["route"] for record in records)
    extension_counts = Counter(record["extension"] or "[none]" for record in records)
    term_counts = Counter(record["term"] or "[other]" for record in records)
    course_counts = Counter(record["course"] or "[program-wide-or-other]" for record in records)
    duplicate_groups = sum(1 for paths_for_hash in by_hash.values() if len(paths_for_hash) > 1)
    report = {
        "generated_at": utc_now(),
        "root": str(ROOT),
        "file_count": len(records),
        "unique_content_count": len(by_hash),
        "duplicate_group_count": duplicate_groups,
        "total_bytes": total_bytes,
        "total_gib": round(total_bytes / (1024**3), 3),
        "review_queue_count": review_count,
        "routes": dict(sorted(route_counts.items())),
        "extensions": dict(sorted(extension_counts.items())),
        "terms": dict(sorted(term_counts.items())),
        "courses": dict(sorted(course_counts.items())),
        "excluded_directories": sorted(IGNORED_DIRS),
    }
    json_dump(RAG_DIR / "inventory-report.json", report)
    print(json.dumps(report, indent=2, ensure_ascii=False))


def text_value(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            pass
    return str(value)


def base_record(manifest_record: dict[str, Any], **extra: Any) -> dict[str, Any]:
    return {
        "document_id": manifest_record["document_id"],
        "artifact_version": manifest_record["content_hash"],
        "source_path": manifest_record["source_path"],
        "source_paths": manifest_record.get("source_paths", [manifest_record["source_path"]]),
        "course": manifest_record.get("course"),
        "term": manifest_record.get("term"),
        "content_role": manifest_record.get("content_role"),
        "access_review": manifest_record.get("access_review"),
        "structured_data": False,
        **extra,
    }


def pdf_page_count(path: Path) -> int | None:
    result = command_output(["pdfinfo", str(path)], timeout=120)
    if result.returncode != 0:
        return None
    match = re.search(r"^Pages:\s+(\d+)\s*$", result.stdout, re.MULTILINE)
    return int(match.group(1)) if match else None


def ocr_pdf_page(path: Path, page_number: int, cache_dir: Path) -> str:
    cache_dir.mkdir(parents=True, exist_ok=True)
    image_prefix = cache_dir / f"page-{page_number}"
    image_path = image_prefix.with_suffix(".png")
    if not image_path.exists():
        rendered = command_output(
            [
                "pdftoppm",
                "-f",
                str(page_number),
                "-l",
                str(page_number),
                "-singlefile",
                "-png",
                "-r",
                "180",
                str(path),
                str(image_prefix),
            ],
            timeout=180,
        )
        if rendered.returncode != 0:
            return ""
    # Automatic page segmentation handles both slide layouts and book pages
    # more reliably than treating the whole page as one uniform text block.
    result = command_output(["tesseract", str(image_path), "stdout", "-l", "eng", "--psm", "3"], timeout=180)
    return result.stdout.strip() if result.returncode == 0 else ""


def extract_pdf(path: Path, record: dict[str, Any], enable_ocr: bool) -> Iterator[dict[str, Any]]:
    page_count = pdf_page_count(path)
    result = command_output(["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"], timeout=600)
    if result.returncode != 0:
        yield base_record(
            record,
            record_type="error",
            locator_type="document",
            locator_value="",
            text="",
            extraction_method="pdftotext",
            error=(result.stderr or "pdftotext failed").strip(),
        )
        return
    pages = result.stdout.split("\f")
    if pages and not pages[-1].strip():
        pages.pop()
    count = page_count or len(pages)
    cache_dir = RAG_DIR / "ocr-cache" / record["content_hash"]
    ocr_pages: dict[int, str] = {}
    candidates = [
        page_number
        for page_number in range(1, count + 1)
        if enable_ocr
        and len(re.sub(r"\s+", "", pages[page_number - 1] if page_number <= len(pages) else "")) < 24
    ]
    if candidates:
        workers = min(max(1, int(os.environ.get("RAG_OCR_WORKERS", "6"))), len(candidates))
        with ThreadPoolExecutor(max_workers=workers) as executor:
            pending = {executor.submit(ocr_pdf_page, path, page_number, cache_dir): page_number for page_number in candidates}
            for future in as_completed(pending):
                page_number = pending[future]
                try:
                    ocr_pages[page_number] = future.result()
                except Exception:
                    ocr_pages[page_number] = ""
    for page_number in range(1, count + 1):
        page_text = pages[page_number - 1] if page_number <= len(pages) else ""
        method = "pdftotext"
        if page_number in ocr_pages:
            ocr_text = ocr_pages[page_number]
            if ocr_text:
                page_text = ocr_text
                method = "tesseract"
        if not page_text.strip():
            page_text = (
                f"Visual-only PDF page: {path.name}\n"
                f"Physical page: {page_number} of {count}\n"
                "No OCR-readable text was detected. Inspect the source page for the chart, diagram, or other visual evidence."
            )
            method = "pdf-visual-catalog"
        yield base_record(
            record,
            record_type="text",
            locator_type="pdf_page",
            locator_value=str(page_number),
            page_number=page_number,
            page_count=count,
            section_title=None,
            char_start=0,
            char_end=len(page_text),
            text=page_text.strip(),
            extraction_method=method,
        )


def shape_text(shape: Any) -> list[str]:
    lines: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            lines.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                lines.append("\t".join(values))
    return lines


def extract_pptx(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))
        return
    try:
        presentation = Presentation(str(path))
        for slide_number, slide in enumerate(presentation.slides, 1):
            lines: list[str] = []
            for shape in slide.shapes:
                lines.extend(shape_text(shape))
            try:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    lines.append(f"Speaker notes: {notes_text}")
            except Exception:
                pass
            text = "\n".join(line for line in lines if line.strip())
            yield base_record(
                record,
                record_type="text",
                locator_type="pptx_slide",
                locator_value=str(slide_number),
                slide_number=slide_number,
                slide_count=len(presentation.slides),
                section_title=None,
                char_start=0,
                char_end=len(text),
                text=text,
                extraction_method="python-pptx",
            )
    except Exception as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))


def extract_docx(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))
        return
    try:
        document = Document(str(path))
        blocks: list[tuple[str | None, str]] = []
        section_title: str | None = None
        for paragraph_number, paragraph in enumerate(document.paragraphs, 1):
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
            if style_name.lower().startswith("heading"):
                section_title = text
            blocks.append((section_title, f"{text}"))
        for table_number, table in enumerate(document.tables, 1):
            rows = []
            for row in table.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                rows.append("\t".join(values))
            table_text = "\n".join(row for row in rows if row.strip())
            if table_text:
                blocks.append((section_title, f"Table {table_number}\n{table_text}"))

        text = "\n\n".join(block for _, block in blocks)
        for index, start in enumerate(range(0, len(text), TEXT_CHUNK_CHARS)):
            part = text[start : start + TEXT_CHUNK_CHARS].strip()
            if not part:
                continue
            nearest_section = next((title for title, block in reversed(blocks) if block and block in part), section_title)
            yield base_record(
                record,
                record_type="text",
                locator_type="docx_block",
                locator_value=str(index + 1),
                section_title=nearest_section,
                char_start=start,
                char_end=start + len(part),
                text=part,
                extraction_method="python-docx",
            )
    except Exception as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))


def converted_office_path(path: Path, target_suffix: str, temp_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    result = command_output(
        [
            soffice,
            f"-env:UserInstallation={(temp_dir / 'profile').resolve().as_uri()}",
            "--headless",
            "--convert-to",
            target_suffix.lstrip("."),
            "--outdir",
            str(temp_dir),
            str(path),
        ],
        timeout=600,
    )
    if result.returncode != 0:
        return None
    candidate = temp_dir / f"{path.stem}.{target_suffix.lstrip('.') }"
    return candidate if candidate.exists() else None


def extract_legacy_document(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    with tempfile.TemporaryDirectory(prefix="rag-doc-") as temp_name:
        converted = converted_office_path(path, ".docx", Path(temp_name))
        if converted is None:
            yield base_record(
                record,
                record_type="error",
                locator_type="document",
                locator_value="",
                text="",
                extraction_method="libreoffice",
                error="Could not convert legacy .doc with LibreOffice",
            )
            return
        yield from extract_docx(converted, record)


def extract_spreadsheet(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    source_path = path
    temp_context: tempfile.TemporaryDirectory[str] | None = None
    if path.suffix.lower() == ".xls":
        temp_context = tempfile.TemporaryDirectory(prefix="rag-xls-")
        converted = converted_office_path(path, ".xlsx", Path(temp_context.name))
        if converted is None:
            yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error="Could not convert legacy .xls with LibreOffice")
            temp_context.cleanup()
            return
        source_path = converted
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))
        if temp_context:
            temp_context.cleanup()
        return
    try:
        workbook = load_workbook(filename=str(source_path), read_only=True, data_only=False, keep_links=False)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            # openpyxl exposes chartsheets in workbook.sheetnames, but they
            # do not contain cell rows. Keep them in the workbook metadata and
            # skip them here rather than turning a valid workbook into a
            # partial extraction failure.
            if not hasattr(sheet, "iter_rows"):
                continue
            batch: list[str] = []
            batch_start: int | None = None
            batch_end: int | None = None
            max_col = 1

            def emit_batch() -> dict[str, Any] | None:
                nonlocal batch, batch_start, batch_end, max_col
                if not batch or batch_start is None or batch_end is None:
                    return None
                first_row = batch_start
                last_row = batch_end
                last_col = get_column_letter(max_col)
                locator = f"{sheet_name}!A{first_row}:{last_col}{last_row}"
                text = f"Workbook: {path.name}\nSheet: {sheet_name}\nRange: {locator}\n" + "\n".join(batch)
                batch = []
                batch_start = None
                batch_end = None
                max_col = 1
                return base_record(
                    record,
                    record_type="text",
                    locator_type="spreadsheet_range",
                    locator_value=locator,
                    sheet_name=sheet_name,
                    range_start=f"A{first_row}",
                    range_end=f"{last_col}{last_row}",
                    char_start=0,
                    char_end=len(text),
                    text=text,
                    structured_data=True,
                    extraction_method="openpyxl",
                )

            for row_number, row in enumerate(sheet.iter_rows(values_only=False), 1):
                values = [text_value(cell.value).replace("\n", " ") for cell in row]
                while values and not values[-1]:
                    values.pop()
                if not values:
                    continue
                max_col = max(max_col, len(values))
                line = "\t".join(values)
                if batch_start is None:
                    batch_start = row_number
                batch_end = row_number
                batch.append(f"{row_number}: {line}")
                if len(batch) >= 40 or sum(len(item) for item in batch) >= 12000:
                    emitted = emit_batch()
                    if emitted:
                        yield emitted
            emitted = emit_batch()
            if emitted:
                yield emitted
        workbook.close()
    except Exception as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))
    finally:
        if temp_context:
            temp_context.cleanup()


def extract_dataset(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    columns: list[str] = []
    sample_rows: list[list[str]] = []
    row_count = 0
    try:
        with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            for row in reader:
                row_count += 1
                if row_count == 1:
                    columns = row[:100]
                elif len(sample_rows) < 10:
                    sample_rows.append(row[:100])
    except Exception as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))
        return
    sample = "\n".join("\t".join(value for value in row) for row in sample_rows)
    text = (
        f"Dataset: {path.name}\n"
        f"Format: {path.suffix.lower().lstrip('.')}\n"
        f"Approximate rows including header: {row_count}\n"
        f"Columns: {', '.join(columns)}\n"
        f"Sample rows:\n{sample}"
    )
    yield base_record(
        record,
        record_type="dataset_catalog",
        locator_type="dataset_schema",
        locator_value="schema",
        char_start=0,
        char_end=len(text),
        text=text,
        structured_data=True,
        row_count=row_count,
        columns=columns,
        extraction_method="csv-schema-and-sample",
        full_row_embedding=False,
    )


def extract_notebook(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    try:
        notebook = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        cells = notebook.get("cells", [])
        for cell_number, cell in enumerate(cells, 1):
            text = "".join(cell.get("source", []))
            if not text.strip():
                continue
            cell_type = cell.get("cell_type", "unknown")
            yield base_record(
                record,
                record_type="text",
                locator_type="notebook_cell",
                locator_value=str(cell_number),
                cell_number=cell_number,
                cell_type=cell_type,
                char_start=0,
                char_end=len(text),
                text=text.strip(),
                extraction_method="nbformat-json",
            )
    except Exception as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))


def extract_text_file(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))
        return
    lines = text.splitlines()
    if not lines:
        return
    section_title: str | None = None
    block_size = 180
    for start_line in range(0, len(lines), block_size):
        end_line = min(len(lines), start_line + block_size)
        block = "\n".join(lines[start_line:end_line]).strip()
        if not block:
            continue
        headings = [line.strip().lstrip("#").strip() for line in lines[start_line:end_line] if line.lstrip().startswith("#")]
        if headings:
            section_title = headings[-1]
        yield base_record(
            record,
            record_type="text",
            locator_type="line_range",
            locator_value=f"{start_line + 1}-{end_line}",
            line_start=start_line + 1,
            line_end=end_line,
            section_title=section_title,
            char_start=sum(len(line) + 1 for line in lines[:start_line]),
            char_end=sum(len(line) + 1 for line in lines[:end_line]),
            text=block,
            extraction_method="utf8-lines",
        )


def extract_image(path: Path, record: dict[str, Any], enable_ocr: bool) -> Iterator[dict[str, Any]]:
    width = height = None
    frame_count = None
    try:
        from PIL import Image

        with Image.open(path) as image:
            width, height = image.size
            frame_count = getattr(image, "n_frames", 1)
    except Exception:
        pass
    ocr_text = ""
    method = "image-metadata"
    if enable_ocr and shutil.which("tesseract"):
        result = command_output(["tesseract", str(path), "stdout", "-l", "eng", "--psm", "6"], timeout=180)
        if result.returncode == 0:
            ocr_text = result.stdout.strip()
            method = "image-metadata-and-tesseract"
    text = f"Image: {path.name}\nDimensions: {width}x{height}\nFrames: {frame_count}\n"
    if ocr_text:
        text += f"OCR text:\n{ocr_text}"
    yield base_record(
        record,
        record_type="image_catalog",
        locator_type="image_asset",
        locator_value="asset",
        char_start=0,
        char_end=len(text),
        text=text,
        width=width,
        height=height,
        frame_count=frame_count,
        extraction_method=method,
    )


def extract_archive(path: Path, record: dict[str, Any]) -> Iterator[dict[str, Any]]:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            lines = [f"Archive: {path.name}", f"Members: {len(members)}"]
            for member in members[:500]:
                lines.append(f"{member.filename}\t{member.file_size} bytes")
            if len(members) > 500:
                lines.append(f"... {len(members) - 500} additional members not listed")
            text = "\n".join(lines)
            yield base_record(
                record,
                record_type="archive_catalog",
                locator_type="archive_member_catalog",
                locator_value="members",
                char_start=0,
                char_end=len(text),
                text=text,
                member_count=len(members),
                structured_data=True,
                extraction_method="zipfile-catalog",
            )
    except Exception as exc:
        yield base_record(record, record_type="error", locator_type="document", locator_value="", text="", error=str(exc))


def extract_records(path: Path, record: dict[str, Any], enable_ocr: bool) -> Iterator[dict[str, Any]]:
    route = record["route"]
    if route == "pdf":
        yield from extract_pdf(path, record, enable_ocr)
    elif route == "pptx":
        yield from extract_pptx(path, record)
    elif route == "docx":
        yield from extract_docx(path, record)
    elif route == "legacy_document":
        yield from extract_legacy_document(path, record)
    elif route in {"spreadsheet", "legacy_spreadsheet"}:
        yield from extract_spreadsheet(path, record)
    elif route == "dataset":
        yield from extract_dataset(path, record)
    elif route == "notebook":
        yield from extract_notebook(path, record)
    elif route in {"source_code", "text", "markdown"}:
        yield from extract_text_file(path, record)
    elif route == "image":
        yield from extract_image(path, record, enable_ocr)
    elif route == "archive":
        yield from extract_archive(path, record)
    else:
        yield base_record(
            record,
            record_type="unsupported",
            locator_type="document",
            locator_value="",
            text="",
            extraction_method="none",
            error=f"No extractor for route {route}",
        )


def extract(args: argparse.Namespace) -> None:
    manifest_path = RAG_DIR / "manifest.jsonl"
    if not manifest_path.exists():
        raise SystemExit("Missing .rag/manifest.jsonl. Run inventory first.")
    EXTRACTED_DIR.mkdir(parents=True, exist_ok=True)
    records = list(jsonl_read(manifest_path))
    canonical = [record for record in records if not record.get("duplicate_of")]
    summary: dict[str, Any] = {
        "generated_at": utc_now(),
        "source_manifest": str(manifest_path),
        "requested_documents": len(canonical),
        "ocr_enabled": bool(args.ocr),
        "documents": [],
    }
    for index, record in enumerate(canonical, 1):
        path = ROOT / record["source_path"]
        output_path = EXTRACTED_DIR / f"{record['content_hash']}.jsonl"
        if output_path.exists() and not args.force:
            cached_rows = list(jsonl_read(output_path))
            errors = sum(1 for row in cached_rows if row.get("record_type") == "error")
            cached_status = "cached_with_errors" if errors else "cached"
            summary["documents"].append({"source_path": record["source_path"], "status": cached_status, "records": len(cached_rows), "errors": errors})
            continue
        try:
            extracted_rows = list(extract_records(path, record, args.ocr))
            jsonl_write(output_path, extracted_rows)
            errors = sum(1 for row in extracted_rows if row.get("record_type") == "error")
            status = "error" if errors and len(extracted_rows) == errors else "complete_with_errors" if errors else "complete"
            summary["documents"].append({"source_path": record["source_path"], "status": status, "records": len(extracted_rows), "errors": errors})
        except Exception as exc:
            error_row = base_record(record, record_type="error", locator_type="document", locator_value="", text="", extraction_method="pipeline", error=str(exc))
            jsonl_write(output_path, [error_row])
            summary["documents"].append({"source_path": record["source_path"], "status": "error", "records": 1, "errors": 1, "error": str(exc)})
        print(f"extract: {index}/{len(canonical)} {record['source_path']}", file=sys.stderr)
    summary["completed"] = sum(1 for item in summary["documents"] if item["status"] in {"complete", "cached"})
    summary["failed"] = sum(1 for item in summary["documents"] if item["status"] in {"error", "complete_with_errors", "cached_with_errors"})
    json_dump(RAG_DIR / "extraction-report.json", summary)
    print(json.dumps({k: v for k, v in summary.items() if k != "documents"}, indent=2))


def normalize_whitespace(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_text(text: str, max_chars: int = TEXT_CHUNK_CHARS, overlap: int = TEXT_CHUNK_OVERLAP) -> Iterator[tuple[int, int, str]]:
    if len(text) <= max_chars:
        yield 0, len(text), text
        return
    start = 0
    while start < len(text):
        hard_end = min(len(text), start + max_chars)
        end = hard_end
        if hard_end < len(text):
            boundaries = [text.rfind("\n\n", start + max_chars // 2, hard_end), text.rfind("\n", start + max_chars // 2, hard_end), text.rfind(". ", start + max_chars // 2, hard_end)]
            best = max(boundaries)
            if best > start:
                end = best + (2 if text[best : best + 2] == ". " else 1 if text[best : best + 1] == "\n" else 0)
        segment = text[start:end].strip()
        if segment:
            leading = len(text[start:end]) - len(text[start:end].lstrip())
            trailing = len(text[start:end]) - len(text[start:end].rstrip())
            yield start + leading, end - trailing, segment
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)


def normalize(args: argparse.Namespace) -> None:
    if not EXTRACTED_DIR.exists():
        raise SystemExit("Missing .rag/extracted. Run extract first.")
    manifest_path = RAG_DIR / "manifest.jsonl"
    if not manifest_path.exists():
        raise SystemExit("Missing .rag/manifest.jsonl. Run inventory first.")
    active_hashes = {
        record["content_hash"]
        for record in jsonl_read(manifest_path)
        if not record.get("duplicate_of")
    }
    output_path = RAG_DIR / "chunks.jsonl"
    chunk_count = 0
    record_count = 0
    error_count = 0
    documents: Counter[str] = Counter()
    locators: Counter[str] = Counter()

    def rows() -> Iterator[dict[str, Any]]:
        nonlocal chunk_count, record_count, error_count
        for extracted_path in sorted(EXTRACTED_DIR.glob("*.jsonl")):
            # A changed source receives a new content-hash filename. Ignore
            # old derived artifacts so a refresh cannot reintroduce stale
            # chunks from a previous version of the archive.
            if extracted_path.stem not in active_hashes:
                continue
            for source_record in jsonl_read(extracted_path):
                record_count += 1
                if source_record.get("record_type") == "error":
                    error_count += 1
                    continue
                raw_text = normalize_whitespace(source_record.get("text", ""))
                if not raw_text:
                    continue
                source_start = int(source_record.get("char_start") or 0)
                parent_id = hashlib.sha1(
                    f"{source_record.get('document_id')}|{source_record.get('locator_type')}|{source_record.get('locator_value')}".encode()
                ).hexdigest()
                pieces = list(split_text(raw_text))
                for part_number, (start, end, text) in enumerate(pieces, 1):
                    chunk_key = f"{source_record.get('document_id')}|{source_record.get('locator_type')}|{source_record.get('locator_value')}|{part_number}|{text}"
                    chunk_id = hashlib.sha256(chunk_key.encode("utf-8")).hexdigest()
                    chunk = {
                        "chunk_id": f"chunk:{chunk_id}",
                        "parent_id": f"parent:{parent_id}",
                        "document_id": source_record.get("document_id"),
                        "artifact_version": source_record.get("artifact_version"),
                        "source_path": source_record.get("source_path"),
                        "source_paths": source_record.get("source_paths", []),
                        "course": source_record.get("course"),
                        "term": source_record.get("term"),
                        "content_role": source_record.get("content_role"),
                        "access_review": source_record.get("access_review"),
                        "record_type": source_record.get("record_type"),
                        "locator_type": source_record.get("locator_type"),
                        "locator_value": source_record.get("locator_value"),
                        "page_number": source_record.get("page_number"),
                        "slide_number": source_record.get("slide_number"),
                        "sheet_name": source_record.get("sheet_name"),
                        "section_title": source_record.get("section_title"),
                        "line_start": source_record.get("line_start"),
                        "line_end": source_record.get("line_end"),
                        "char_start": source_start + start,
                        "char_end": source_start + end,
                        "part_number": part_number,
                        "part_count": len(pieces),
                        "token_estimate": max(1, round(len(text) / 4)),
                        "structured_data": bool(source_record.get("structured_data")),
                        "text": text,
                        "extraction_method": source_record.get("extraction_method"),
                    }
                    for field in ('archive_members', 'member_hash', 'member_locator', 'ocr_words',
                                  'ocr_confidence', 'width', 'height', 'frame_count', 'extraction_warning'):
                        if field in source_record:
                            chunk[field] = source_record[field]
                    chunk_count += 1
                    documents[str(chunk.get("document_id"))] += 1
                    locators[str(chunk.get("locator_type"))] += 1
                    yield chunk

    jsonl_write(output_path, rows())
    report = {
        "generated_at": utc_now(),
        "source_directory": str(EXTRACTED_DIR),
        "chunk_count": chunk_count,
        "source_record_count": record_count,
        "source_error_count": error_count,
        "document_count": len(documents),
        "chunks_by_locator": dict(sorted(locators.items())),
        "chunks_by_document_top_20": documents.most_common(20),
        "chunk_chars": TEXT_CHUNK_CHARS,
        "chunk_overlap_chars": TEXT_CHUNK_OVERLAP,
    }
    json_dump(RAG_DIR / "normalization-report.json", report)
    print(json.dumps(report, indent=2, ensure_ascii=False))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Stages 1-3 for the Course Archive RAG")
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("inventory", help="Create .rag/manifest.jsonl and inventory reports")
    extract_parser = subparsers.add_parser("extract", help="Extract supported files with source locators")
    extract_parser.add_argument("--ocr", action="store_true", help="OCR low-text PDF pages and image assets")
    extract_parser.add_argument("--force", action="store_true", help="Re-extract files even if cached output exists")
    subparsers.add_parser("normalize", help="Create canonical .rag/chunks.jsonl")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    if args.command == "inventory":
        inventory()
    elif args.command == "extract":
        extract(args)
    elif args.command == "normalize":
        normalize(args)


if __name__ == "__main__":
    main()
